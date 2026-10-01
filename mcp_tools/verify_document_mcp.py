"""Exercise real stdio discovery, six native formats, local OCR, and write guards."""
from __future__ import annotations

import asyncio
import json
import os
from pathlib import Path
import sys
import tempfile

from mcp import Client
from mcp.client.stdio import StdioServerParameters


def fixtures(root: Path) -> list[Path]:
    from docx import Document
    from openpyxl import Workbook
    from PIL import Image, ImageDraw, ImageFont
    from pptx import Presentation
    from pptx.util import Inches
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject

    (root / "sample.txt").write_text("Native text 12345", encoding="utf-8")
    (root / "sample.csv").write_text("name,value\nalpha,12345\n", encoding="utf-8")
    workbook = Workbook()
    workbook.active.append(["name", "value"])
    workbook.active.append(["alpha", 12345])
    workbook.save(root / "sample.xlsx")
    document = Document()
    document.add_paragraph("Native DOCX 12345")
    document.save(root / "sample.docx")
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text = "Native PPTX 12345"
    deck.save(root / "sample.pptx")
    writer = PdfWriter()
    page = writer.add_blank_page(width=600, height=300)
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 24 Tf 40 150 Td (Native PDF 12345) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    writer.write(root / "sample.pdf")
    image = Image.new("RGB", (1300, 260), "white")
    ImageDraw.Draw(image).text((40, 80), "LOCAL OCR CHECK 12345", font=ImageFont.load_default(size=70), fill="black")
    image.save(root / "sample.png")
    image.save(root / "scanned.pdf", resolution=144)
    embedded = Document()
    embedded.add_paragraph("Native body")
    embedded.add_picture(str(root / "sample.png"))
    embedded.save(root / "embedded.docx")
    capped = PdfWriter()
    capped.append(root / "scanned.pdf")
    capped.append(root / "scanned.pdf")
    capped.write(root / "capped.pdf")
    (root / "guard.md").write_text("# Test\n\n## Current\nold\n\n## Retained\nkeep\n", encoding="utf-8")
    return [root / ("sample" + extension) for extension in (".txt", ".csv", ".xlsx", ".docx", ".pptx", ".pdf")]


async def verify() -> dict:
    with tempfile.TemporaryDirectory(prefix="document-mcp-verify-") as directory:
        root = Path(directory)
        sources = fixtures(root)
        script = Path(__file__).with_name("document_server.py")
        params = StdioServerParameters(command=sys.executable, args=["-B", str(script), "--read-root", str(root), "--write-root", str(root)], env={**os.environ, "PYTHONUTF8": "1"})
        async with Client(params, read_timeout_seconds=180) as client:
            tools = (await client.list_tools()).tools
            names = sorted(tool.name for tool in tools)
            assert names == ["document_status", "extract_document", "inspect_markdown", "update_markdown"], names
            async def call(name, arguments):
                result = await client.call_tool(name, arguments)
                assert not result.is_error, result.content
                assert isinstance(result.structured_content, dict)
                return result.structured_content
            status = await call("document_status", {})
            assert status["status"] == "ready", status
            for source in sources:
                result = await call("extract_document", {"source": str(source), "ocr": "off"})
                assert result["written"] is False
                assert "12345" in result["content"], source
            guard = root / "guard.md"
            before = guard.read_bytes()
            inspected = await call("inspect_markdown", {"target": str(guard)})
            arguments = {"target": str(guard), "content": "## Current\nnew", "expected_sha256": inspected["sha256"], "heading": "## Current"}
            preview = await call("update_markdown", arguments)
            assert preview["status"] == "preview" and guard.read_bytes() == before
            applied = await call("update_markdown", {**arguments, "write": True})
            assert applied["status"] == "written" and "## Retained\nkeep" in guard.read_text()
            stale = await client.call_tool("update_markdown", {**arguments, "write": True})
            assert stale.is_error
            outside = await client.call_tool("inspect_markdown", {"target": str(root.parent / "outside.md")})
            assert outside.is_error
            bad_bool = await client.call_tool("update_markdown", {**arguments, "write": "false"})
            assert bad_bool.is_error
            output = root / "audit.md"
            written = await call("extract_document", {"source": str(sources[0]), "output": str(output), "write_output": True, "ocr": "off"})
            assert written["written"] and output.is_file()
            conflict = await client.call_tool("extract_document", {"source": str(sources[0]), "output": str(output), "write_output": True, "ocr": "off"})
            assert conflict.is_error
            ocr_results = []
            native = await call("extract_document", {"source": str(root / "sample.pdf")})
            assert native["ocr"]["status"] == "not_needed"
            for file in ("sample.png", "scanned.pdf", "embedded.docx"):
                result = await call("extract_document", {"source": str(root / file), "ocr_max_items": 1})
                assert result["ocr"]["status"] == "used", result["ocr"]
                text = result["ocr"]["records"][0]["text"]
                assert "LOCAL" in text and "12345" in text, text
                ocr_results.append({"format": Path(file).suffix, "status": "passed", "text": text})
            capped = await call("extract_document", {"source": str(root / "capped.pdf"), "ocr_max_items": 1})
            assert capped["status"] == "partial" and capped["ocr"]["omitted_items"] == 1
            selected = await call("extract_document", {"source": str(root / "capped.pdf"), "pages": [2], "ocr_max_items": 1})
            assert selected["ocr"]["records"][0]["location"] == "PDF page 2"
            return {"status": "passed", "protocol": client.protocol_version, "tools": names,
                    "native_formats": [path.suffix for path in sources], "ocr": ocr_results,
                    "guards": ["preview", "hash", "section_preservation", "path_roots", "strict_boolean", "output_overwrite", "ocr_auto_native_skip", "ocr_item_cap", "ocr_page_selection"],
                    "versions": status["versions"]}


if __name__ == "__main__":
    print(json.dumps(asyncio.run(asyncio.wait_for(verify(), timeout=300)), ensure_ascii=False, indent=2))
