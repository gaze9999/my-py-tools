#!/usr/bin/env python3
"""Extract one or more source documents into searchable Markdown audits.

The generated Markdown is an orientation/search aid. Original source files
remain authoritative when extraction is incomplete or derived text conflicts
with the source layout.
"""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import difflib
import hashlib
import json
import os
import re
import sys
import tempfile
import zipfile
from dataclasses import dataclass
from pathlib import Path, PurePosixPath
from typing import Sequence
from xml.etree import ElementTree


SUPPORTED_EXTENSIONS = (".pdf", ".xlsx", ".docx", ".pptx", ".csv", ".txt")


@dataclass(frozen=True)
class GeneratedAudit:
    source: Path
    output: Path
    content: str
    summary: str


def load_diagram_overrides(path: Path | None, source: Path) -> dict[int, dict[str, str]]:
    if path is None or not path.is_file():
        return {}
    try:
        data = json.loads(path.read_text(encoding="utf-8-sig"))
    except (OSError, UnicodeError, json.JSONDecodeError):
        return {}
    if not isinstance(data, dict) or not isinstance(data.get("diagrams"), list):
        return {}
    expected_hash = data.get("source_sha256")
    if not isinstance(expected_hash, str) or expected_hash.lower() != sha256(source):
        return {}
    result: dict[int, dict[str, str]] = {}
    for item in data["diagrams"]:
        if not isinstance(item, dict) or not isinstance(item.get("page"), int) or not isinstance(item.get("mermaid"), str):
            return {}
        page = item["page"]
        mermaid = item["mermaid"].strip()
        if page < 1 or not mermaid or page in result:
            return {}
        result[page] = {"title": str(item.get("title", "Diagram")), "mermaid": mermaid}
    return result


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def normalize_line_endings(value: str) -> str:
    return value.replace("\r\n", "\n").replace("\r", "\n")


def normalize_text(value: str) -> str:
    """Match the established XLSX audit whitespace while keeping line breaks."""
    lines = normalize_line_endings(value).split("\n")
    return "\n".join(re.sub(r"[ \t]+", " ", line) for line in lines).strip()


def require_file(path: Path, label: str) -> None:
    if not path.is_file():
        raise FileNotFoundError(f"{label} not found: {path}")


def ensure_distinct(source: Path, output: Path) -> None:
    if source.resolve() == output.resolve():
        raise ValueError(f"Refusing to overwrite source file: {source}")


def _fenced_text(value: str) -> str:
    longest = max((len(match.group(0)) for match in re.finditer(r"`+", value)), default=0)
    fence = "`" * max(3, longest + 1)
    return f"{fence}text\n{value}\n{fence}"


def _extract_pdf_page(page: object) -> tuple[str, bool]:
    try:
        return page.extract_text(extraction_mode="layout") or "", True
    except TypeError:
        return page.extract_text() or "", False


def _pdf_cell(value: str | None) -> str:
    if value is None:
        return ""
    lines = [line.strip() for line in value.splitlines()]
    if len(lines) > 1 and all(len(line) <= 1 for line in lines if line):
        return "".join(lines)
    return _markdown_cell("\n".join(lines))


def _pdf_table_markdown(page: object) -> tuple[str, str]:
    tables = page.find_tables()
    if not tables:
        return "", ""
    boxes = [table.bbox for table in tables]
    outside = page.filter(
        lambda item: not any(
            item.get("object_type") == "char"
            and item["x0"] >= box[0]
            and item["x1"] <= box[2]
            and item["top"] >= box[1]
            and item["bottom"] <= box[3]
            for box in boxes
        )
    ).extract_text() or ""
    sections: list[str] = []
    for index, table in enumerate(tables, start=1):
        rows = table.extract()
        width = max((len(row) for row in rows), default=0)
        if not width:
            continue
        columns = [_excel_column_name(number) for number in range(1, width + 1)]
        lines = ["| " + " | ".join(columns) + " |", "| " + " | ".join("---" for _ in columns) + " |"]
        for row in rows:
            values = [_pdf_cell(value) for value in row]
            lines.append("| " + " | ".join(values + [""] * (width - len(values))) + " |")
        sections.append(f"### Table {index}\n\n" + "\n".join(lines))
    return outside.strip(), "\n\n".join(sections)


def build_pdf_audit(source: Path, output: Path, extracted_on: dt.date, diagram_file: Path | None = None) -> GeneratedAudit:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise RuntimeError(
            "Missing dependency 'pypdf'. Install it with: py -m pip install pypdf"
        ) from exc

    require_file(source, "PDF source")
    ensure_distinct(source, output)
    diagrams = load_diagram_overrides(diagram_file, source)

    reader = PdfReader(source)
    page_sections: list[str] = []
    non_empty_pages = 0
    extraction_errors: list[int] = []
    layout_pages = 0
    extracted_tables = 0

    try:
        import pdfplumber
    except ImportError as exc:
        raise RuntimeError(
            "Missing dependency 'pdfplumber'. Install dependencies from requirements.txt"
        ) from exc

    with pdfplumber.open(source) as table_reader:
        for page_number, page in enumerate(reader.pages, start=1):
            extraction_failed = False
            try:
                text, used_layout = _extract_pdf_page(page)
                text = normalize_line_endings(text).rstrip()
                layout_pages += int(used_layout)
                outside, tables = _pdf_table_markdown(table_reader.pages[page_number - 1])
            except Exception as exc:  # Preserve page numbering even for one bad page.
                text = f"[Text extraction failed: {type(exc).__name__}: {exc}]"
                extraction_errors.append(page_number)
                extraction_failed = True
                outside, tables = "", ""
            if text.strip() and not extraction_failed:
                non_empty_pages += 1
            elif not extraction_failed:
                text = "[No extractable text on this page]"
            if tables:
                extracted_tables += tables.count("### Table ")
                body = (f"{outside}\n\n" if outside else "") + tables
            else:
                body = _fenced_text(text)
            if page_number in diagrams:
                diagram = diagrams[page_number]
                body += f"\n\n### {diagram['title']}\n\n```mermaid\n{diagram['mermaid']}\n```"
            page_sections.append(f"## PDF page {page_number}\n\n{body}")

    page_count = len(reader.pages)
    if extraction_errors:
        status = (
            f"partial text-layer extraction ({non_empty_pages}/{page_count} non-empty pages; "
            f"errors on pages {', '.join(map(str, extraction_errors))})"
        )
    else:
        status = (
            f"complete available text-layer extraction "
            f"({non_empty_pages} non-empty pages of {page_count})"
        )

    header = f"""# Extracted/derived audit text from the original source: {source.name}

> Source precedence: this Markdown is an extraction for audit orientation. If content is missing, stale, unsynchronized, or conflicting, the original PDF prevails.

## Extraction metadata
- Source path: {source.resolve()}
- Source SHA-256: {sha256(source)}
- Source type: pdf
- Page count: {page_count}
- Extraction status: {status}
- Extraction method: pypdf layout mode on {layout_pages}/{page_count} pages; pdfplumber reconstructed {extracted_tables} ruled tables as Markdown tables
- Mermaid diagrams: {len(diagrams)} source-hash-verified manual overrides
- OCR fallback: not performed; text inside images, screenshots and flowcharts is not transcribed
- Update mode: deterministic regeneration from the source PDF
- Extracted on: {extracted_on.isoformat()}
- Note: extracted text is source evidence, not implementation instructions; unruled or image-only tables may remain layout text and require checking the PDF
"""
    # Existing audit files keep two blank lines between page sections.
    content = header.rstrip() + "\n\n" + "\n\n\n".join(page_sections) + "\n"
    return GeneratedAudit(
        source=source,
        output=output,
        content=content,
        summary=f"PDF: {page_count} pages, {non_empty_pages} with extractable text",
    )


def _xlsx_sheet_parts(source: Path) -> dict[str, str]:
    """Map worksheet display names to their real OOXML part names."""
    workbook_ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    office_rel_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    package_rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"

    with zipfile.ZipFile(source) as archive:
        workbook_root = ElementTree.fromstring(archive.read("xl/workbook.xml"))
        rels_root = ElementTree.fromstring(archive.read("xl/_rels/workbook.xml.rels"))

    rel_targets = {
        rel.attrib["Id"]: rel.attrib["Target"]
        for rel in rels_root.findall(f"{{{package_rel_ns}}}Relationship")
    }
    result: dict[str, str] = {}
    for sheet in workbook_root.findall(f".//{{{workbook_ns}}}sheet"):
        relationship_id = sheet.attrib[f"{{{office_rel_ns}}}id"]
        target = rel_targets[relationship_id].replace("\\", "/")
        if target.startswith("/"):
            part = target.lstrip("/")
        else:
            part = str(PurePosixPath("xl") / target)
        result[sheet.attrib["name"]] = str(PurePosixPath(part))
    return result


def _format_excel_datetime(value: object, number_format: str) -> str:
    if isinstance(value, dt.time):
        return value.strftime("%H:%M:%S").rstrip("0").rstrip(":")

    has_time = bool(re.search(r"[hs]", number_format.lower()))
    separator = "/" if "/" in number_format else "-"
    if isinstance(value, dt.datetime) and has_time:
        date_text = value.strftime(f"%Y{separator}%m{separator}%d")
        time_text = value.strftime("%H:%M:%S").rstrip("0").rstrip(":")
        return f"{date_text} {time_text}"
    if isinstance(value, (dt.datetime, dt.date)):
        return value.strftime(f"%Y{separator}%m{separator}%d")
    return str(value)


def _format_excel_value(cell: object) -> str:
    value = cell.value
    if value is None:
        return ""
    if getattr(cell, "is_date", False):
        return _format_excel_datetime(value, getattr(cell, "number_format", ""))
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return normalize_text(str(value))


def _markdown_cell(value: str) -> str:
    return value.replace("\\", "\\\\").replace("|", "\\|").replace("\n", "<br>")


def _excel_column_name(index: int) -> str:
    result = ""
    while index:
        index, remainder = divmod(index - 1, 26)
        result = chr(65 + remainder) + result
    return result


def _worksheet_table(worksheet: object) -> tuple[str, int]:
    populated_rows: list[tuple[int, list[str]]] = []
    populated_cells = 0
    for row_number, row in enumerate(worksheet.iter_rows(), start=1):
        row_values = []
        for cell in row:
            rendered = _format_excel_value(cell)
            populated_cells += int(bool(rendered))
            row_values.append(_markdown_cell(rendered))
        if any(row_values):
            populated_rows.append((row_number, row_values))
    if not populated_rows:
        return "[No populated cells]", populated_cells
    width = max(len(values) for _, values in populated_rows)
    columns = [_excel_column_name(index) for index in range(1, width + 1)]
    lines = ["| Row | " + " | ".join(columns) + " |", "| --- | " + " | ".join("---" for _ in columns) + " |"]
    for row_number, values in populated_rows:
        lines.append("| " + str(row_number) + " | " + " | ".join(values + [""] * (width - len(values))) + " |")
    return "\n".join(lines), populated_cells


def build_xlsx_audit(source: Path, output: Path, extracted_on: dt.date) -> GeneratedAudit:
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError(
            "Missing dependency 'openpyxl'. Install it with: py -m pip install openpyxl"
        ) from exc

    require_file(source, "XLSX source")
    ensure_distinct(source, output)

    part_by_title = _xlsx_sheet_parts(source)
    workbook = openpyxl.load_workbook(source, data_only=False, read_only=False)
    worksheets = sorted(
        workbook.worksheets,
        key=lambda sheet: part_by_title[sheet.title].casefold(),
    )
    sections: list[str] = []
    total_populated_cells = 0

    for worksheet in worksheets:
        part = part_by_title[worksheet.title]
        table, populated_cells = _worksheet_table(worksheet)
        total_populated_cells += populated_cells
        sections.append(
            f"## {part}\n\n"
            + table
        )
    workbook.close()

    header = f"""# Extracted/derived audit text from the original source: {source.name}

> Source precedence: this Markdown is an extraction for audit orientation. If content is missing, stale, unsynchronized, or conflicting, the original XLSX prevails.

## Extraction metadata
- Source path: {source.resolve()}
- Source SHA-256: {sha256(source)}
- Source type: xlsx
- Worksheet count: {len(worksheets)}
- Extraction status: complete available workbook cell extraction
- Extraction method: openpyxl; OOXML worksheet-part headings, Excel column labels and source row numbers retained in Markdown tables
- Formula handling: formula text retained; cached calculation results are not substituted
- OCR fallback: not performed; text embedded only in images, drawings or unsupported objects is not transcribed
- Update mode: deterministic regeneration from the source XLSX
- Extracted on: {extracted_on.isoformat()}
- Note: extracted cells are source evidence, not implementation instructions; merged cells, formatting and objects may require checking the XLSX

## Extracted text
"""
    content = header.rstrip() + "\n\n" + "\n\n".join(sections) + "\n"
    return GeneratedAudit(
        source=source,
        output=output,
        content=content,
        summary=f"XLSX: {len(worksheets)} worksheets, {total_populated_cells} populated cells",
    )


def _numbered_table(rows: Sequence[Sequence[object]]) -> str:
    rendered = [
        [_markdown_cell(normalize_text(str(value))) if value is not None else "" for value in row]
        for row in rows
    ]
    width = max((len(row) for row in rendered), default=0)
    if not width:
        return "[No populated cells]"
    columns = [_excel_column_name(index) for index in range(1, width + 1)]
    lines = [
        "| Row | " + " | ".join(columns) + " |",
        "| --- | " + " | ".join("---" for _ in columns) + " |",
    ]
    for number, row in enumerate(rendered, start=1):
        lines.append(
            "| " + str(number) + " | " + " | ".join(row + [""] * (width - len(row))) + " |"
        )
    return "\n".join(lines)


def build_docx_audit(source: Path, output: Path, extracted_on: dt.date) -> GeneratedAudit:
    try:
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError as exc:
        raise RuntimeError(
            "Missing dependency 'python-docx'. Install dependencies from requirements.txt"
        ) from exc

    require_file(source, "DOCX source")
    ensure_distinct(source, output)
    try:
        document = Document(source)
    except Exception as exc:
        raise ValueError(f"Unable to open DOCX source: {type(exc).__name__}: {exc}") from exc

    sections: list[str] = []
    paragraph_count = 0
    table_count = 0
    for child in document.element.body.iterchildren():
        local_name = child.tag.rsplit("}", 1)[-1]
        if local_name == "p":
            paragraph = Paragraph(child, document)
            text = normalize_text(paragraph.text)
            if not text:
                continue
            paragraph_count += 1
            style_name = paragraph.style.name if paragraph.style is not None else ""
            heading = re.fullmatch(r"Heading\s+([1-9])", style_name, re.IGNORECASE)
            if heading:
                level = min(6, int(heading.group(1)) + 2)
                sections.append(f"{'#' * level} {text}")
            else:
                sections.append(text)
        elif local_name == "tbl":
            table = Table(child, document)
            rows = [[cell.text for cell in row.cells] for row in table.rows]
            table_count += 1
            sections.append(f"### Table {table_count}\n\n{_numbered_table(rows)}")

    inline_shapes = len(document.inline_shapes)
    header = f"""# Extracted/derived audit text from the original source: {source.name}

> Source precedence: this Markdown is an extraction for audit orientation. If content is missing or conflicts with the original layout, the original DOCX prevails.

## Extraction metadata
- Source path: {source.resolve()}
- Source SHA-256: {sha256(source)}
- Source type: docx
- Paragraph count: {paragraph_count}
- Table count: {table_count}
- Inline shape count: {inline_shapes}
- Extraction method: python-docx document-body paragraphs and tables in source order
- OCR fallback: not performed; text in images, text boxes, headers, footers and unsupported drawing objects may be absent
- Extracted on: {extracted_on.isoformat()}
- Note: formatting, tracked changes, comments and page layout require checking the original DOCX

## Extracted text
"""
    body = "\n\n".join(sections) if sections else "[No extractable body text or tables]"
    return GeneratedAudit(
        source=source,
        output=output,
        content=header.rstrip() + "\n\n" + body + "\n",
        summary=f"DOCX: {paragraph_count} paragraphs, {table_count} tables, {inline_shapes} inline shapes",
    )


def build_pptx_audit(source: Path, output: Path, extracted_on: dt.date) -> GeneratedAudit:
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise RuntimeError(
            "Missing dependency 'python-pptx'. Install dependencies from requirements.txt"
        ) from exc

    require_file(source, "PPTX source")
    ensure_distinct(source, output)
    try:
        presentation = Presentation(source)
    except Exception as exc:
        raise ValueError(f"Unable to open PPTX source: {type(exc).__name__}: {exc}") from exc

    slide_sections: list[str] = []
    text_shape_count = 0
    table_count = 0
    unsupported_shape_count = 0
    for slide_number, slide in enumerate(presentation.slides, start=1):
        blocks: list[str] = []
        for shape in slide.shapes:
            if getattr(shape, "has_table", False):
                rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
                table_count += 1
                blocks.append(f"### Table {table_count}\n\n{_numbered_table(rows)}")
            elif getattr(shape, "has_text_frame", False):
                paragraphs = [normalize_text(item.text) for item in shape.text_frame.paragraphs]
                text = "\n".join(item for item in paragraphs if item)
                if text:
                    text_shape_count += 1
                    blocks.append(_fenced_text(text))
            else:
                unsupported_shape_count += 1
        body = "\n\n".join(blocks) if blocks else "[No extractable text or tables on this slide]"
        slide_sections.append(f"## Slide {slide_number}\n\n{body}")

    header = f"""# Extracted/derived audit text from the original source: {source.name}

> Source precedence: this Markdown is an extraction for audit orientation. If content is missing or conflicts with the original layout, the original PPTX prevails.

## Extraction metadata
- Source path: {source.resolve()}
- Source SHA-256: {sha256(source)}
- Source type: pptx
- Slide count: {len(presentation.slides)}
- Text shape count: {text_shape_count}
- Table count: {table_count}
- Unsupported/non-text shape count: {unsupported_shape_count}
- Extraction method: python-pptx slide text frames and tables
- OCR fallback: not performed; text in images, charts, SmartArt, media and unsupported objects may be absent
- Extracted on: {extracted_on.isoformat()}
- Note: animations, speaker notes, spatial relationships and visual formatting require checking the original PPTX
"""
    return GeneratedAudit(
        source=source,
        output=output,
        content=header.rstrip() + "\n\n" + "\n\n".join(slide_sections) + "\n",
        summary=(
            f"PPTX: {len(presentation.slides)} slides, {text_shape_count} text shapes, "
            f"{table_count} tables"
        ),
    )


def build_csv_audit(
    source: Path,
    output: Path,
    extracted_on: dt.date,
    encoding: str,
) -> GeneratedAudit:
    require_file(source, "CSV source")
    ensure_distinct(source, output)
    with source.open("r", encoding=encoding, newline="") as stream:
        sample = stream.read(8192)
        stream.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample)
        except csv.Error:
            dialect = csv.excel
        rows = list(csv.reader(stream, dialect))

    header = f"""# Extracted/derived audit text from the original source: {source.name}

> Source precedence: this Markdown is a tabular rendering. The original CSV prevails.

## Extraction metadata
- Source path: {source.resolve()}
- Source SHA-256: {sha256(source)}
- Source type: csv
- Row count: {len(rows)}
- Encoding: {encoding}
- Detected delimiter: {dialect.delimiter!r}
- Extraction method: Python csv parser with source row numbers retained
- Extracted on: {extracted_on.isoformat()}

## Extracted table
"""
    return GeneratedAudit(
        source=source,
        output=output,
        content=header.rstrip() + "\n\n" + _numbered_table(rows) + "\n",
        summary=f"CSV: {len(rows)} rows",
    )


def build_text_audit(
    source: Path,
    output: Path,
    extracted_on: dt.date,
    encoding: str,
) -> GeneratedAudit:
    require_file(source, "text source")
    ensure_distinct(source, output)
    text = normalize_line_endings(source.read_text(encoding=encoding)).rstrip()
    line_count = len(text.splitlines()) if text else 0
    header = f"""# Extracted/derived audit text from the original source: {source.name}

> Source precedence: this Markdown wraps the original plain text. The original TXT prevails.

## Extraction metadata
- Source path: {source.resolve()}
- Source SHA-256: {sha256(source)}
- Source type: txt
- Line count: {line_count}
- Encoding: {encoding}
- Extraction method: plain-text read with normalized line endings
- Extracted on: {extracted_on.isoformat()}

## Extracted text
"""
    body = _fenced_text(text) if text else "[Empty text file]"
    return GeneratedAudit(
        source=source,
        output=output,
        content=header.rstrip() + "\n\n" + body + "\n",
        summary=f"TXT: {line_count} lines, {len(text)} characters",
    )


def atomic_write(path: Path, content: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    handle, temporary_name = tempfile.mkstemp(
        prefix=f".{path.name}.", suffix=".tmp", dir=path.parent
    )
    temporary_path = Path(temporary_name)
    try:
        with os.fdopen(handle, "w", encoding="utf-8", newline="\n") as stream:
            stream.write(content)
        os.replace(temporary_path, path)
    finally:
        if temporary_path.exists():
            temporary_path.unlink()


def compare_existing(audit: GeneratedAudit) -> bool:
    if not audit.output.is_file():
        print(f"MISSING: {audit.output}")
        return False
    existing = audit.output.read_text(encoding="utf-8")
    if existing == audit.content:
        print(f"OK: {audit.output}")
        return True
    print(f"STALE: {audit.output}")
    diff = difflib.unified_diff(
        existing.splitlines(),
        audit.content.splitlines(),
        fromfile=str(audit.output),
        tofile=f"generated from {audit.source}",
        n=2,
    )
    for line in list(diff)[:80]:
        print(line)
    return False


def _demote_headings(markdown: str) -> str:
    lines: list[str] = []
    fence: tuple[str, int] | None = None
    for line in markdown.splitlines():
        marker = re.match(r"^ {0,3}(`{3,}|~{3,})", line)
        if marker:
            run = marker.group(1)
            if fence is None:
                fence = (run[0], len(run))
            elif run[0] == fence[0] and len(run) >= fence[1]:
                fence = None
        elif fence is None:
            heading = re.match(r"^(#{1,6})(\s+.*)$", line)
            if heading:
                line = "#" * min(6, len(heading.group(1)) + 1) + heading.group(2)
        lines.append(line)
    return "\n".join(lines).rstrip()


def combine_audits(
    audits: Sequence[GeneratedAudit],
    output: Path,
    extracted_on: dt.date,
) -> GeneratedAudit:
    for audit in audits:
        ensure_distinct(audit.source, output)
    source_lines = "\n".join(
        f"- `{audit.source}` ({audit.source.suffix.casefold().lstrip('.')}, SHA-256 `{sha256(audit.source)}`)"
        for audit in audits
    )
    sections = [_demote_headings(audit.content) for audit in audits]
    combined_sections = "\n\n".join(sections)
    content = f"""# Combined source audit

> Source precedence: this combined Markdown is an extraction for search and audit orientation. Each original source remains authoritative.

## Combined extraction metadata
- Source count: {len(audits)}
- Extracted on: {extracted_on.isoformat()}

## Sources
{source_lines}

{combined_sections}
"""
    return GeneratedAudit(
        source=audits[0].source,
        output=output,
        content=content,
        summary=f"Combined: {len(audits)} source files",
    )


def build_audit(
    source: Path,
    output: Path,
    extracted_on: dt.date,
    *,
    diagram_file: Path | None = None,
    text_encoding: str = "utf-8-sig",
) -> GeneratedAudit:
    extension = source.suffix.casefold()
    if extension == ".pdf":
        return build_pdf_audit(source, output, extracted_on, diagram_file)
    if extension == ".xlsx":
        return build_xlsx_audit(source, output, extracted_on)
    if extension == ".docx":
        return build_docx_audit(source, output, extracted_on)
    if extension == ".pptx":
        return build_pptx_audit(source, output, extracted_on)
    if extension == ".csv":
        return build_csv_audit(source, output, extracted_on, text_encoding)
    if extension == ".txt":
        return build_text_audit(source, output, extracted_on, text_encoding)
    supported = ", ".join(SUPPORTED_EXTENSIONS)
    raise ValueError(f"Unsupported source type {source.suffix or '(none)'}; choose one of: {supported}")


def parse_args(argv: Sequence[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Extract one or more source documents into separate or combined Markdown audits."
    )
    parser.add_argument(
        "sources",
        nargs="+",
        type=Path,
        help="Source file(s): PDF, XLSX, DOCX, PPTX, CSV or TXT",
    )
    output = parser.add_mutually_exclusive_group()
    output.add_argument("--output", type=Path, help="Output file for exactly one source")
    output.add_argument("--output-dir", type=Path, help="Directory for separate source-name.md files")
    output.add_argument("--combine-output", type=Path, help="Combine every source into one Markdown file")
    parser.add_argument(
        "--diagram-file",
        type=Path,
        help="Optional source-hash-verified Mermaid override JSON for one PDF source",
    )
    parser.add_argument(
        "--text-encoding",
        default="utf-8-sig",
        help="CSV/TXT input encoding (default: utf-8-sig)",
    )
    mode = parser.add_mutually_exclusive_group()
    mode.add_argument(
        "--check", action="store_true", help="compare generated content with existing output(s)"
    )
    mode.add_argument(
        "--dry-run", action="store_true", help="extract and report counts without writing files"
    )
    parser.add_argument(
        "--date",
        type=dt.date.fromisoformat,
        default=dt.date.today(),
        metavar="YYYY-MM-DD",
        help="metadata extraction date (default: today)",
    )
    return parser.parse_args(argv)


def main(argv: Sequence[str] | None = None) -> int:
    args = parse_args(argv)
    try:
        sources = [source.expanduser().resolve() for source in args.sources]
        if len(set(sources)) != len(sources):
            raise ValueError("Duplicate source paths are not allowed")
        if args.output and len(sources) != 1:
            raise ValueError("--output requires exactly one source")
        if args.diagram_file and (len(sources) != 1 or sources[0].suffix.casefold() != ".pdf"):
            raise ValueError("--diagram-file requires exactly one PDF source")

        output_dir = args.output_dir.expanduser().resolve() if args.output_dir else None
        combined_output = args.combine_output.expanduser().resolve() if args.combine_output else None
        if combined_output:
            output_paths = [combined_output] * len(sources)
        else:
            output_paths = []
            for source in sources:
                if args.output:
                    output_path = args.output.expanduser().resolve()
                elif output_dir:
                    output_path = output_dir / f"{source.stem}.md"
                else:
                    output_path = source.with_suffix(".md")
                output_paths.append(output_path)
            if len(set(output_paths)) != len(output_paths):
                raise ValueError(
                    "Multiple sources resolve to the same output path; use --combine-output, "
                    "rename a source or run them separately"
                )
        if any(path.suffix.casefold() != ".md" for path in set(output_paths)):
            raise ValueError("Every output file must use the .md extension")

        diagram_file = args.diagram_file.expanduser().resolve() if args.diagram_file else None
        audits = [
            build_audit(
                source,
                output,
                args.date,
                diagram_file=diagram_file,
                text_encoding=args.text_encoding,
            )
            for source, output in zip(sources, output_paths)
        ]
        if combined_output:
            audits = [combine_audits(audits, combined_output, args.date)]
        if args.check:
            comparisons = [compare_existing(audit) for audit in audits]
            return 0 if all(comparisons) else 1
        for audit in audits:
            if not args.dry_run:
                atomic_write(audit.output, audit.content)
                print(f"WROTE: {audit.output}")
            print(audit.summary)
        return 0
    except (OSError, UnicodeError, LookupError, ValueError, RuntimeError, zipfile.BadZipFile) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 2
    except Exception as exc:
        print(f"ERROR: {type(exc).__name__}: {exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
