from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path

from audit.source_audit_extract import load_diagram_overrides, main, sha256


class SourceAuditExtractTests(unittest.TestCase):
    def test_single_text_input_defaults_to_same_directory(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory, "notes.txt")
            source.write_text("first\nsecond\n", encoding="utf-8")

            self.assertEqual(main([str(source), "--date", "2026-09-22"]), 0)

            output = source.with_suffix(".md")
            content = output.read_text(encoding="utf-8")
            self.assertIn("Source type: txt", content)
            self.assertIn("Extracted on: 2026-09-22 00:00:00", content)
            self.assertIn("first\nsecond", content)
            self.assertEqual(
                main([str(source), "--check", "--date", "2026-09-22"]),
                0,
            )
            output.write_text(content + "stale\n", encoding="utf-8")
            self.assertEqual(
                main([str(source), "--check", "--date", "2026-09-22"]),
                1,
            )

    def test_extracted_at_preserves_hours_minutes_and_seconds(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory, "notes.txt")
            output = Path(directory, "audit.md")
            source.write_text("timestamp", encoding="utf-8")

            self.assertEqual(
                main(
                    [
                        str(source),
                        "--output",
                        str(output),
                        "--extracted-at",
                        "2026-09-22T14:35:27",
                    ]
                ),
                0,
            )
            self.assertIn(
                "Extracted on: 2026-09-22 14:35:27",
                output.read_text(encoding="utf-8"),
            )

    def test_multiple_inputs_support_separate_and_combined_outputs(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            text_source = root / "notes.txt"
            csv_source = root / "rows.csv"
            text_source.write_text("hello", encoding="utf-8")
            csv_source.write_text("name,value\nalpha,1\n", encoding="utf-8")
            output_dir = root / "separate"

            self.assertEqual(
                main([str(text_source), str(csv_source), "--output-dir", str(output_dir)]),
                0,
            )
            self.assertTrue((output_dir / "notes.md").is_file())
            self.assertTrue((output_dir / "rows.md").is_file())

            combined = root / "combined.md"
            self.assertEqual(
                main([str(text_source), str(csv_source), "--combine-output", str(combined)]),
                0,
            )
            content = combined.read_text(encoding="utf-8")
            self.assertIn("# Combined source audit", content)
            self.assertIn("notes.txt", content)
            self.assertIn("rows.csv", content)

    def test_office_inputs_generate_markdown(self) -> None:
        import openpyxl
        from docx import Document
        from pptx import Presentation

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            xlsx = root / "book.xlsx"
            docx = root / "document.docx"
            pptx = root / "slides.pptx"

            workbook = openpyxl.Workbook()
            workbook.active.append(["field", "value"])
            workbook.active.append(["alpha", 1])
            workbook.save(xlsx)
            workbook.close()

            document = Document()
            document.add_heading("Heading", level=1)
            document.add_paragraph("Paragraph")
            table = document.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "field"
            table.cell(0, 1).text = "value"
            table.cell(1, 0).text = "alpha"
            table.cell(1, 1).text = "1"
            document.save(docx)

            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[1])
            slide.shapes.title.text = "Title"
            slide.placeholders[1].text = "Body"
            presentation.save(pptx)

            output_dir = root / "output"
            self.assertEqual(
                main([str(xlsx), str(docx), str(pptx), "--output-dir", str(output_dir)]),
                0,
            )
            for name in ("book.md", "document.md", "slides.md"):
                self.assertTrue((output_dir / name).is_file(), name)

    def test_output_option_rejects_multiple_sources(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            first = root / "first.txt"
            second = root / "second.txt"
            first.write_text("one", encoding="utf-8")
            second.write_text("two", encoding="utf-8")

            self.assertEqual(main([str(first), str(second), "--output", str(root / "out.md")]), 2)

    def test_diagram_override_is_optional_and_hash_guarded(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "source.pdf"
            source.write_bytes(b"fixture")
            override = root / "diagrams.json"

            self.assertEqual(load_diagram_overrides(root / "missing.json", source), {})
            override.write_text("not-json", encoding="utf-8")
            self.assertEqual(load_diagram_overrides(override, source), {})
            override.write_text(
                json.dumps({"diagrams": [{"page": 1, "mermaid": "flowchart LR\nA-->B"}]}),
                encoding="utf-8",
            )
            self.assertEqual(load_diagram_overrides(override, source), {})
            override.write_text(
                json.dumps(
                    {
                        "source_sha256": sha256(source),
                        "diagrams": [{"page": 1, "title": "Flow", "mermaid": "flowchart LR\nA-->B"}],
                    }
                ),
                encoding="utf-8",
            )
            self.assertEqual(load_diagram_overrides(override, source)[1]["title"], "Flow")


if __name__ == "__main__":
    unittest.main()
